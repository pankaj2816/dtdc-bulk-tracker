import os
import io
import re
import json
import base64
import time
from datetime import datetime
import urllib.request
import urllib.parse
from collections import defaultdict
from flask import Flask, render_template, request, jsonify, send_file, Response
from flask_cors import CORS
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import tempfile
try:
    import ddddocr
except Exception:
    ddddocr = None

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
app = Flask(
    __name__, 
    template_folder=os.path.join(BASE_DIR, 'templates'), 
    static_folder=os.path.join(BASE_DIR, 'static')
)
app.config['TEMPLATES_AUTO_RELOAD'] = True
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max file size
CORS(app)

class VercelPathMiddleware:
    def __init__(self, wsgi_app):
        self.wsgi_app = wsgi_app

    def __call__(self, environ, start_response):
        query = environ.get('QUERY_STRING', '')
        if '__path__=' in query:
            import urllib.parse
            parsed = urllib.parse.parse_qs(query)
            if '__path__' in parsed and parsed['__path__']:
                raw_path = parsed['__path__'][0]
                if not raw_path.startswith('/'):
                    raw_path = '/' + raw_path
                while '//' in raw_path:
                    raw_path = raw_path.replace('//', '/')
                environ['PATH_INFO'] = raw_path
                remaining = {k: v for k, v in parsed.items() if k != '__path__'}
                environ['QUERY_STRING'] = urllib.parse.urlencode(remaining, doseq=True)
        return self.wsgi_app(environ, start_response)

app.wsgi_app = VercelPathMiddleware(app.wsgi_app)

UPLOAD_FOLDER = os.path.join(tempfile.gettempdir(), 'dtdc_uploads')
try:
    os.makedirs(UPLOAD_FOLDER, exist_ok=True)
except Exception:
    pass

USER_AGENT = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'

_ocr_instance = None

def get_ocr_engine():
    global _ocr_instance
    if _ocr_instance is None:
        try:
            import ddddocr
            _ocr_instance = ddddocr.DdddOcr(show_ad=False)
            print("OCR Engine Ready!", flush=True)
        except Exception as e:
            print("Warning: OCR Engine failed to initialize:", e, flush=True)
    return _ocr_instance



def fetch_dtdc_captcha():
    """Fetch a fresh captcha from DTDC."""
    url = f"https://www.dtdc.com/wp-json/custom/v1/generate-captcha?t={int(os.times().elapsed * 1000)}"
    req = urllib.request.Request(
        url,
        headers={
            'User-Agent': USER_AGENT,
            'Referer': 'https://www.dtdc.com/track-your-shipment/'
        }
    )
    with urllib.request.urlopen(req, timeout=15) as resp:
        data = json.loads(resp.read().decode('utf-8'))
        return {
            'key': data.get('key'),
            'image': data.get('image')
        }


def validate_dtdc_captcha(key, value):
    """Validate captcha and obtain DTDC Track Token."""
    url = "https://www.dtdc.com/wp-json/custom/v1/captcha/validate"
    payload = json.dumps({"captchaKey": key, "captchaValue": value}).encode('utf-8')
    req = urllib.request.Request(
        url,
        data=payload,
        headers={
            'Content-Type': 'application/json',
            'User-Agent': USER_AGENT,
            'Referer': 'https://www.dtdc.com/track-your-shipment/'
        }
    )
    with urllib.request.urlopen(req, timeout=15) as resp:
        return json.loads(resp.read().decode('utf-8'))


def get_verified_token_auto(max_attempts=8):
    """Automatically fetch, OCR-solve, and validate captcha with DTDC."""
    for attempt in range(max_attempts):
        try:
            c = fetch_dtdc_captcha()
            img_bytes = base64.b64decode(c['image'])
            ocr_engine = get_ocr_engine()
            if not ocr_engine:
                print("OCR engine not available", flush=True)
                break
            pred = ocr_engine.classification(img_bytes)
            cleaned = pred.upper().replace(' ', '')
            
            # Validate with DTDC
            val_res = validate_dtdc_captcha(c['key'], cleaned)
            if val_res.get('success') and val_res.get('token'):
                return val_res.get('token')
            time.sleep(0.2)
        except Exception:
            pass
    return None


def pull_dtdc_details(numbers, token):
    """Fetch encrypted payload for list of tracking numbers."""
    url = "https://www.dtdc.com/wp-json/custom/v1/tracking/pull-details"
    payload = json.dumps({"trackNumbers": numbers}).encode('utf-8')
    headers = {
        'Content-Type': 'application/json',
        'User-Agent': USER_AGENT,
        'Referer': 'https://www.dtdc.com/track-your-shipment/'
    }
    if token:
        headers['X-DTDC-Track-Token'] = token

    req = urllib.request.Request(url, data=payload, headers=headers)
    with urllib.request.urlopen(req, timeout=30) as resp:
        return json.loads(resp.read().decode('utf-8'))


def post_to_trackshipment(redirect_url, payload_d):
    """POST payload to trackshipment and retrieve rendered HTML."""
    data = urllib.parse.urlencode({'d': payload_d}).encode('utf-8')
    req = urllib.request.Request(
        redirect_url,
        data=data,
        headers={
            'User-Agent': USER_AGENT,
            'Referer': 'https://www.dtdc.com/track-your-shipment/',
            'Content-Type': 'application/x-www-form-urlencoded'
        }
    )
    with urllib.request.urlopen(req, timeout=35) as resp:
        return resp.read().decode('utf-8', errors='ignore')


def parse_dtdc_html(html_content, requested_numbers=None):
    """Parse DTDC HTML tracking page and extract structured tracking details."""
    soup = BeautifulSoup(html_content, 'html.parser')
    results = []

    # CASE 1: Multiple shipments view (accordion items)
    cards = soup.find_all('div', class_=re.compile(r'accordion-item'))
    
    if cards:
        for card in cards:
            item = {
                "awb": "",
                "reference_no": "",
                "edd": "",
                "status": "",
                "status_type": "pending",
                "origin": "",
                "destination": "",
                "latest_update": "",
                "latest_time": ""
            }

            header = card.find('div', class_=re.compile(r'accordion-header|accord-row'))
            if header:
                labels = header.find_all('p', class_='label')
                for lbl in labels:
                    txt = lbl.get_text(strip=True).lower()
                    val_el = lbl.find_next_sibling('p', class_='value')
                    val = val_el.get_text(strip=True) if val_el else ""

                    if 'shipment' in txt or 'awb' in txt or 'consignment' in txt:
                        item['awb'] = val
                    elif 'reference' in txt:
                        item['reference_no'] = val
                    elif 'edd' in txt:
                        item['edd'] = val

                status_el = header.find('span', class_=re.compile(r'status'))
                if status_el:
                    item['status'] = re.sub(r'\s+', ' ', status_el.get_text(strip=True))
                    classes = status_el.get('class', [])
                    for c in classes:
                        if c != 'status':
                            item['status_type'] = c.lower()

            body = card.find('div', class_=re.compile(r'accordion-body'))
            if body:
                loc_items = body.find_all('div', class_='location-item')
                for loc in loc_items:
                    lbl = loc.find('span', class_='location-label')
                    val = loc.find('span', class_='location-value')
                    if lbl and val:
                        l_txt = lbl.get_text(strip=True).lower()
                        v_txt = re.sub(r'\s+', ' ', val.get_text(strip=True))
                        if 'origin' in l_txt:
                            item['origin'] = v_txt
                        elif 'destination' in l_txt:
                            item['destination'] = v_txt

                tl_first = body.find('div', class_=re.compile(r'timeline-item'))
                if tl_first:
                    time_span = tl_first.find('span', class_=re.compile(r'timeline-date'))
                    desc_div = tl_first.find('div', class_=re.compile(r'timeline-description'))
                    status_span = tl_first.find('span', class_=re.compile(r'timeline-status'))
                    if time_span:
                        item['latest_time'] = time_span.get_text(strip=True)
                    if desc_div:
                        item['latest_update'] = desc_div.get_text(strip=True)
                    elif status_span:
                        item['latest_update'] = status_span.get_text(strip=True)

            if item['awb']:
                results.append(item)

    # CASE 2: Single shipment view
    if not results:
        awb_sec = soup.find('div', class_=re.compile(r'awb-section'))
        awb_val = ""
        if awb_sec:
            m = re.search(r'([A-Za-z0-9]{8,15})', awb_sec.get_text(strip=True))
            if m:
                awb_val = m.group(1)

        if not awb_val and requested_numbers:
            awb_val = requested_numbers[0]

        if awb_val:
            item = {
                "awb": awb_val,
                "reference_no": "-",
                "edd": "-",
                "status": "",
                "status_type": "pending",
                "origin": "-",
                "destination": "-",
                "latest_update": "-",
                "latest_time": "-"
            }

            ref_el = soup.find(class_=re.compile(r'ref-number|refNumber'))
            if ref_el:
                item['reference_no'] = ref_el.get_text(strip=True)

            badge = soup.select_one('div.status-badge:not(.status-badge-main)')
            if not badge:
                badge = soup.find('span', class_=re.compile(r'timeline-status'))
            if badge:
                item['status'] = re.sub(r'\s+', ' ', badge.get_text(strip=True))

            edd_title = soup.find(class_=re.compile(r'main-status-title|edd-date'))
            if edd_title:
                edd_text = edd_title.get_text(strip=True)
                m_edd = re.search(r'(?:Delivery by|Delivered on)\s*(.*)', edd_text, re.I)
                if m_edd:
                    item['edd'] = m_edd.group(1).strip()
                else:
                    item['edd'] = edd_text

            loc_items = soup.find_all('div', class_=re.compile(r'location-item'))
            for loc in loc_items:
                lbl = loc.find(class_=re.compile(r'location-label'))
                val = loc.find(class_=re.compile(r'location-value'))
                if lbl and val:
                    l_txt = lbl.get_text(strip=True).lower()
                    v_txt = re.sub(r'\s+', ' ', val.get_text(strip=True))
                    if 'origin' in l_txt:
                        item['origin'] = v_txt
                    elif 'destination' in l_txt:
                        item['destination'] = v_txt

            timeline = soup.find(class_=re.compile(r'timeline'))
            if timeline:
                tl_first = timeline.find('div', class_=re.compile(r'timeline-item'))
                if tl_first:
                    time_span = tl_first.find('span', class_=re.compile(r'timeline-date'))
                    desc_div = tl_first.find('div', class_=re.compile(r'timeline-description'))
                    status_span = tl_first.find('span', class_=re.compile(r'timeline-status'))
                    if time_span:
                        item['latest_time'] = time_span.get_text(strip=True)
                    if desc_div:
                        item['latest_update'] = desc_div.get_text(strip=True)
                    elif status_span:
                        item['latest_update'] = status_span.get_text(strip=True)

            if item['status'] or item['origin'] != '-':
                results.append(item)

    # Categorize all parsed items
    for item in results:
        st_lower = item['status'].lower()
        if any(k in st_lower for k in ['undelivered', 'not delivered', 'un-delivered', 'failed', 'cancelled', 'rto', 'return', 'exception']):
            item['category'] = 'Issue/RTO'
        elif 'out for delivery' in st_lower:
            item['category'] = 'Out for Delivery'
        elif 'delivered' in st_lower:
            item['category'] = 'Delivered'
        elif any(k in st_lower for k in ['in transit', 'dispatched', 'arrived', 'transit', 'reached', 'on the way']):
            item['category'] = 'In Transit'
        elif any(k in st_lower for k in ['booked', 'picked up', 'manifest']):
            item['category'] = 'Booked'
        else:
            item['category'] = 'In Transit' if item['status'] else 'Not Found'

    if requested_numbers:
        found_awbs = {r['awb'].upper() for r in results}
        for num in requested_numbers:
            if num.upper() not in found_awbs:
                results.append({
                    "awb": num,
                    "reference_no": "-",
                    "edd": "-",
                    "status": "No record / Invalid AWB",
                    "status_type": "not_found",
                    "category": "Not Found",
                    "origin": "-",
                    "destination": "-",
                    "latest_update": "No tracking record found on DTDC",
                    "latest_time": "-"
                })

    return results


@app.after_request
def add_header(response):
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, post-check=0, pre-check=0, max-age=0'
    response.headers['Pragma'] = 'no-cache'
    response.headers['Expires'] = '-1'
    return response


@app.route('/api/health')
def health():
    return jsonify({
        'status': 'ok',
        'app': 'DTDC Bulk Tracker',
        'cwd': os.getcwd(),
        'template_exists': os.path.exists(os.path.join(BASE_DIR, 'templates', 'index.html'))
    })


@app.route('/')
@app.route('/index')
@app.route('/api/index')
@app.route('/api/index.py')
def index():
    return render_template('index.html')


@app.errorhandler(404)
def handle_404(e):
    if not request.path.startswith('/api/'):
        return render_template('index.html'), 200
    return jsonify({
        'error': 'Not found',
        'path': request.path,
        'PATH_INFO': request.environ.get('PATH_INFO'),
        'x-matched-path': request.headers.get('x-matched-path')
    }), 404


@app.route('/api/captcha', methods=['GET'])
def get_captcha():
    try:
        data = fetch_dtdc_captcha()
        return jsonify({'success': True, 'key': data['key'], 'image': data['image']})
    except Exception as e:
        return jsonify({'success': False, 'error': f'Failed to fetch captcha: {str(e)}'}), 500


def track_trackon_consignments(numbers):
    """Query Trackon multi-tracking endpoint (https://trackon.in/courier-tracking-Multi) for a batch of numbers."""
    import ssl
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE

    query_str = ', '.join(numbers)
    data = urllib.parse.urlencode({
        'awbMultiTrackingId': query_str,
        'btnMulAwbTrack': 'Track'
    }).encode('utf-8')

    req = urllib.request.Request(
        'https://trackon.in/courier-tracking-Multi',
        data=data,
        headers={
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
            'Content-Type': 'application/x-www-form-urlencoded',
            'Referer': 'https://trackon.in/'
        }
    )

    try:
        with urllib.request.urlopen(req, context=ctx, timeout=25) as resp:
            html = resp.read().decode('utf-8', errors='ignore')
    except Exception as e:
        print(f"Trackon query error: {e}", flush=True)
        html = ''

    soup = BeautifulSoup(html, 'html.parser')
    results = []

    def categorize_event(evt):
        u = evt.upper()
        if any(k in u for k in ['UNDELIVERED', 'NOT DELIVERED', 'UN-DELIVERED', 'FAILED', 'CANCEL', 'DAMAGE', 'HOLD', 'REFUSED', 'REJECTED', 'RTO', 'RETURN']):
            return 'Issue/RTO', f'Issue: {evt[:30]}'
        if any(k in u for k in ['OUT FOR DELIVERY', 'OUT FOR DELV', 'RUNSHEET GENERATED', 'OUT FOR RUNSHEET']):
            return 'Out for Delivery', 'Out for Delivery'
        if 'DELIVERED' in u and not any(k in u for k in ['TO HUB', 'TO BRANCH', 'TO AIRPORT', 'BAG']):
            return 'Delivered', 'Delivered'
        if any(k in u for k in ['IN TRANSIT', 'DISPATCHED', 'VEHICLE OUT', 'PROCESSING', 'ARRIVED', 'BOOKED', 'PICKED UP', 'MANIFEST']):
            return 'In Transit', evt
        return 'In Transit', evt

    for num in numbers:
        # Check if explicitly Not Found
        not_found_pattern = rf'Consignment\s*No:\s*{num}\s*\((?:Not\s*Found|No\s*Record)\)'
        if re.search(not_found_pattern, html, re.I):
            results.append({
                'awb': num,
                'category': 'Not Found',
                'status': 'Not Found on Trackon',
                'edd': '-',
                'origin': '-',
                'destination': '-',
                'latest_update': 'No record found',
                'latest_time': '-'
            })
            continue

        c_header = soup.find(string=re.compile(rf'Consignment\s*No:\s*{num}', re.I))
        if not c_header:
            results.append({
                'awb': num,
                'category': 'Not Found',
                'status': 'Not Found / Expired (>75 days)',
                'edd': '-',
                'origin': '-',
                'destination': '-',
                'latest_update': '-',
                'latest_time': '-'
            })
            continue

        header_text = str(c_header)
        edd = '-'
        m_due = re.search(r'DueDate\s*:\s*([0-9/\.\-]+)', header_text, re.I)
        if m_due:
            edd = m_due.group(1).strip()

        tbl = c_header.find_parent('div').find_next('table') if c_header.find_parent('div') else None
        if not tbl:
            results.append({
                'awb': num,
                'category': 'Booked',
                'status': 'Booked / Processing',
                'edd': edd,
                'origin': '-',
                'destination': '-',
                'latest_update': 'Details pending',
                'latest_time': '-'
            })
            continue

        rows = tbl.find_all('tr')
        if len(rows) <= 1:
            results.append({
                'awb': num,
                'category': 'Booked',
                'status': 'Booked / Processing',
                'edd': edd,
                'origin': '-',
                'destination': '-',
                'latest_update': 'No scan events yet',
                'latest_time': '-'
            })
            continue

        # Row 1 is latest event
        cells_latest = [c.get_text(strip=True) for c in rows[1].find_all(['td', 'th'])]
        latest_date = cells_latest[0] if len(cells_latest) > 0 else '-'
        latest_loc = cells_latest[2] if len(cells_latest) > 2 else '-'
        latest_evt = cells_latest[4] if len(cells_latest) > 4 else (cells_latest[3] if len(cells_latest) > 3 else '-')

        # Earliest event (last row)
        cells_earliest = [c.get_text(strip=True) for c in rows[-1].find_all(['td', 'th'])]
        origin_loc = cells_earliest[2] if len(cells_earliest) > 2 else '-'
        booking_date = cells_earliest[0] if len(cells_earliest) > 0 else '-'

        cat, status_text = categorize_event(latest_evt)

        results.append({
            'awb': num,
            'category': cat,
            'status': status_text,
            'edd': edd,
            'origin': origin_loc,
            'destination': latest_loc,
            'booking_date': booking_date,
            'latest_update': f'{latest_loc}: {latest_evt}',
            'latest_time': latest_date
        })

    return results


@app.route('/api/track-stream', methods=['POST'])
def track_stream():
    """Stream live tracking results in safe 25-number chunks with automated captcha solving, supporting DTDC, Trackon, or Both simultaneously."""
    data = request.json or {}
    raw_numbers = data.get('numbers', [])
    req_dtdc = data.get('dtdc_numbers')
    req_trackon = data.get('trackon_numbers')
    courier = data.get('courier', 'both').lower()

    dtdc_nums = []
    trackon_nums = []

    if req_dtdc is not None or req_trackon is not None:
        for n in (req_dtdc or []):
            c = re.sub(r'[^a-zA-Z0-9]', '', str(n).strip()).upper()
            if c and c not in dtdc_nums:
                dtdc_nums.append(c)
        for n in (req_trackon or []):
            c = re.sub(r'[^a-zA-Z0-9]', '', str(n).strip())
            if c and c not in trackon_nums:
                trackon_nums.append(c)
    else:
        for n in raw_numbers:
            cleaned = re.sub(r'[^a-zA-Z0-9]', '', str(n).strip())
            if not cleaned:
                continue
            if courier == 'dtdc':
                if cleaned.upper() not in dtdc_nums:
                    dtdc_nums.append(cleaned.upper())
            elif courier == 'trackon':
                if cleaned not in trackon_nums:
                    trackon_nums.append(cleaned)
            else:  # courier == 'both'
                if re.match(r'^[0-9]{10,12}$', cleaned):
                    if cleaned not in trackon_nums:
                        trackon_nums.append(cleaned)
                else:
                    if cleaned.upper() not in dtdc_nums:
                        dtdc_nums.append(cleaned.upper())

    total = len(trackon_nums) + len(dtdc_nums)
    if total == 0:
        return jsonify({'success': False, 'error': 'No valid tracking numbers provided.'}), 400

    def generate_events():
        CHUNK_SIZE = 25
        completed = 0
        all_accumulated = []

        # Phase 1: Track Trackon consignments (if any)
        if trackon_nums:
            total_trackon_batches = (len(trackon_nums) + CHUNK_SIZE - 1) // CHUNK_SIZE
            for i in range(0, len(trackon_nums), CHUNK_SIZE):
                chunk = trackon_nums[i:i + CHUNK_SIZE]
                batch_num = (i // CHUNK_SIZE) + 1

                yield f"data: {json.dumps({'type': 'status', 'message': f'Tracking Trackon batch {batch_num} of {total_trackon_batches} ({len(chunk)} parcels)...', 'completed': completed, 'total': total})}\n\n"
                try:
                    chunk_results = track_trackon_consignments(chunk)
                except Exception as e:
                    print(f"Error tracking Trackon chunk: {e}")
                    chunk_results = [{'awb': n, 'category': 'Issue/RTO', 'status': f'Error: {str(e)[:30]}', 'edd': '-', 'origin': '-', 'destination': '-', 'latest_update': '-', 'latest_time': '-'} for n in chunk]

                for r in chunk_results:
                    r['courier'] = 'trackon'
                    r['sheet'] = 'TRACKON'

                completed += len(chunk)
                all_accumulated.extend(chunk_results)

                payload = {
                    'type': 'progress',
                    'completed': completed,
                    'total': total,
                    'percent': int((completed / total) * 100),
                    'new_results': chunk_results
                }
                yield f"data: {json.dumps(payload)}\n\n"
                time.sleep(0.2)

        # Phase 2: Track DTDC consignments (if any)
        if dtdc_nums:
            total_dtdc_batches = (len(dtdc_nums) + CHUNK_SIZE - 1) // CHUNK_SIZE
            for i in range(0, len(dtdc_nums), CHUNK_SIZE):
                chunk = dtdc_nums[i:i + CHUNK_SIZE]
                batch_num = (i // CHUNK_SIZE) + 1

                yield f"data: {json.dumps({'type': 'status', 'message': f'Verifying DTDC security token for batch {batch_num} of {total_dtdc_batches}...', 'completed': completed, 'total': total})}\n\n"

                token = get_verified_token_auto(max_attempts=8)
                if not token:
                    yield f"data: {json.dumps({'type': 'error', 'message': f'Could not verify security token for batch {batch_num}. Retrying...'})}\n\n"
                    token = get_verified_token_auto(max_attempts=8)
                    if not token:
                        continue

                yield f"data: {json.dumps({'type': 'status', 'message': f'Tracking DTDC batch {batch_num} of {total_dtdc_batches} ({len(chunk)} parcels)...', 'completed': completed, 'total': total})}\n\n"

                try:
                    pull_res = pull_dtdc_details(chunk, token)
                    if pull_res.get('success') and pull_res.get('redirect') and pull_res.get('payload'):
                        html = post_to_trackshipment(pull_res['redirect'], pull_res['payload'])
                        chunk_results = parse_dtdc_html(html, requested_numbers=chunk)
                    else:
                        chunk_results = parse_dtdc_html("", requested_numbers=chunk)
                except Exception as e:
                    print(f"Error tracking DTDC chunk: {e}")
                    chunk_results = parse_dtdc_html("", requested_numbers=chunk)

                for r in chunk_results:
                    r['courier'] = 'dtdc'
                    r['sheet'] = 'MIA-2'

                completed += len(chunk)
                all_accumulated.extend(chunk_results)

                payload = {
                    'type': 'progress',
                    'completed': completed,
                    'total': total,
                    'percent': int((completed / total) * 100),
                    'new_results': chunk_results
                }
                yield f"data: {json.dumps(payload)}\n\n"
                time.sleep(0.2)

        # Finished all batches
        yield f"data: {json.dumps({'type': 'done', 'total': total, 'count': len(all_accumulated)})}\n\n"

    return Response(generate_events(), mimetype='text/event-stream')


def is_valid_date_str(s):
    if not s or not isinstance(s, str):
        return False
    m = re.match(r'^(\d{1,2})[\./](\d{1,2})[\./](\d{2,4})$', s.strip())
    if not m:
        return False
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 100:
        y += 2000
    if mo < 1 or mo > 12 or d < 1 or d > 31:
        return False
    now = datetime.now()
    if y < 2024 or y > now.year:
        return False
    try:
        ts = datetime(y, mo, d).timestamp()
        if ts > now.timestamp() + 86400:
            return False
    except Exception:
        return False
    return True


def parse_dmy(s):
    if not s or not isinstance(s, str):
        return 0
    parts = [int(p) for p in re.split(r'[./]', s) if p.isdigit()]
    if len(parts) < 3:
        return 0
    d, mo, y = parts[0], parts[1], parts[2]
    if y < 100:
        y += 2000
    try:
        return datetime(y, mo, d).timestamp()
    except Exception:
        return 0


def extract_sheet_dockets(wb, sheet_target, courier_type):
    target_sheet = None
    for s in wb.sheetnames:
        if s.strip().upper() == sheet_target.upper():
            target_sheet = s
            break
    if not target_sheet:
        for s in wb.sheetnames:
            if sheet_target.upper() in s.strip().upper():
                target_sheet = s
                break
    if not target_sheet:
        return None

    ws = wb[target_sheet]
    date_dockets = defaultdict(list)
    all_dockets = []
    curr_date = "Unknown"

    for row in ws.iter_rows(values_only=True):
        row_date = None
        row_dockets = []
        for cell in row:
            if not cell:
                continue
            c_str = str(cell).strip()
            if 'SHYAM' in c_str.upper():
                continue
            if re.match(r'^\d{1,2}[\.\/]\d{1,2}[\.\/]\d{2,4}$', c_str) and is_valid_date_str(c_str):
                row_date = c_str
            elif re.search(r'\b(\d{1,2}\.\d{1,2}\.\d{2,4})\b', c_str):
                m = re.search(r'\b(\d{1,2}\.\d{1,2}\.\d{2,4})\b', c_str)
                if m and is_valid_date_str(m.group(1)):
                    row_date = m.group(1)
            
            if courier_type == 'dtdc':
                if re.match(r'^[A-Z]{1,4}[0-9]{6,12}$', c_str, re.I):
                    row_dockets.append(c_str.upper())
            else:  # trackon
                if re.match(r'^[0-9]{10,12}$', c_str):
                    row_dockets.append(c_str)

        if row_date:
            curr_date = row_date

        for d in row_dockets:
            if d not in all_dockets:
                all_dockets.append(d)
            if d not in date_dockets[curr_date]:
                date_dockets[curr_date].append(d)

    valid_dates = [d for d in date_dockets.keys() if d != "Unknown" and is_valid_date_str(d)]
    valid_dates.sort(key=parse_dmy)
    latest_date = valid_dates[-1] if valid_dates else None

    date_options = []
    if latest_date:
        date_options.append({
            'label': f"Latest Date: {latest_date} ({len(date_dockets[latest_date])} dockets)",
            'value': latest_date,
            'count': len(date_dockets[latest_date]),
            'numbers': date_dockets[latest_date]
        })

    if len(valid_dates) >= 2:
        last_3 = valid_dates[-3:]
        comb_3 = []
        for d in last_3:
            comb_3.extend(date_dockets[d])
        date_options.append({
            'label': f"Last 3 Days ({len(comb_3)} dockets)",
            'value': 'last_3',
            'count': len(comb_3),
            'numbers': comb_3
        })

    if len(valid_dates) >= 4:
        last_7 = valid_dates[-7:]
        comb_7 = []
        for d in last_7:
            comb_7.extend(date_dockets[d])
        date_options.append({
            'label': f"Last 7 Days ({len(comb_7)} dockets)",
            'value': 'last_7',
            'count': len(comb_7),
            'numbers': comb_7
        })

    for d in reversed(valid_dates[-7:]):
        if d != latest_date:
            date_options.append({
                'label': f"Date {d} ({len(date_dockets[d])} dockets)",
                'value': d,
                'count': len(date_dockets[d]),
                'numbers': date_dockets[d]
            })

    default_numbers = date_options[0]['numbers'] if date_options else all_dockets[:50]

    return {
        'sheet': target_sheet,
        'total_dockets': len(all_dockets),
        'latest_date': latest_date,
        'date_options': date_options,
        'numbers': default_numbers,
        'count': len(default_numbers),
        'all_numbers': all_dockets
    }


@app.route('/api/upload-excel', methods=['POST'])
def upload_excel():
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': 'No file uploaded.'}), 400

    file = request.files['file']
    if not file.filename:
        return jsonify({'success': False, 'error': 'Empty filename.'}), 400

    filename = file.filename
    ext = os.path.splitext(filename)[1].lower()

    if ext not in ['.xlsx', '.xls', '.csv']:
        return jsonify({'success': False, 'error': 'Please upload an Excel (.xlsx, .xls) or CSV file.'}), 400

    try:
        temp_path = os.path.join(UPLOAD_FOLDER, filename)
        file.save(temp_path)

        wb = openpyxl.load_workbook(temp_path, read_only=True)
        sheet_names = wb.sheetnames

        user_courier = request.form.get('courier', 'dtdc').lower()
        dtdc_info = extract_sheet_dockets(wb, 'MIA-2', 'dtdc')
        trackon_info = extract_sheet_dockets(wb, 'TRACKON', 'trackon')
        wb.close()

        if not dtdc_info and not trackon_info:
            return jsonify({
                'success': False,
                'error': f"Neither 'MIA-2' (DTDC) nor 'TRACKON' (Trackon) sheets were found in the uploaded workbook. Available sheets: {', '.join(sheet_names[:5])}..."
            }), 400

        # Choose active courier
        if user_courier == 'trackon':
            active_courier = 'trackon' if trackon_info else 'dtdc'
        else:
            active_courier = 'dtdc' if dtdc_info else 'trackon'

        active_info = trackon_info if active_courier == 'trackon' else dtdc_info

        return jsonify({
            'success': True,
            'filename': filename,
            'sheets': sheet_names,
            'active_courier': active_courier,
            'couriers': {
                'dtdc': dtdc_info,
                'trackon': trackon_info
            },
            'chosen_sheet': active_info['sheet'],
            'total_dockets_in_sheet': active_info['total_dockets'],
            'date_options': active_info['date_options'],
            'numbers': active_info['numbers'],
            'count': active_info['count'],
            'latest_date': active_info['latest_date']
        })

    except Exception as e:
        return jsonify({'success': False, 'error': f'Failed to parse file: {str(e)}'}), 500


@app.route('/api/export-excel', methods=['POST'])
def export_excel():
    data = request.json or {}
    results = data.get('results', [])
    mode = data.get('mode', 'summary')  # 'summary' (fast instant) or 'merged' (merge original file)
    courier = data.get('courier', 'dtdc').lower()

    if not results:
        return jsonify({'success': False, 'error': 'No data to export.'}), 400

    status_map = {str(r['awb']).strip().upper(): r for r in results}

    if mode == 'merged':
        files_dir = UPLOAD_FOLDER if (os.path.exists(UPLOAD_FOLDER) and os.listdir(UPLOAD_FOLDER)) else 'uploads'
        original_files = [f for f in os.listdir(files_dir) if 'daily entry' in f.lower() or 'autosaved' in f.lower()] if os.path.exists(files_dir) else []
        base_file = None
        if original_files:
            base_file = os.path.join(files_dir, original_files[-1])
        elif os.path.exists("DAILY ENTRY 2026.xlsx"):
            base_file = "DAILY ENTRY 2026.xlsx"

        if base_file and os.path.exists(base_file):
            try:
                print(f"Merging live status into {base_file} for {courier.upper()}...", flush=True)
                wb = openpyxl.load_workbook(base_file)
                header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                status_font = Font(name="Calibri", size=10, bold=True)
                regular_font = Font(name="Calibri", size=10)

                fill_delivered = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                font_delivered = Font(name="Calibri", size=10, bold=True, color="006100")
                fill_transit = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
                fill_ofd = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                fill_issue = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

                # Target sheets to update: both TRACKON and MIA-2 if courier == 'both'
                targets = []
                if courier in ['trackon', 'both']:
                    targets.append(('TRACKON', 'trackon'))
                if courier in ['dtdc', 'both']:
                    targets.append(('MIA-2', 'dtdc'))

                for target_sheet_name, c_type in targets:
                    target_sheets = [s for s in wb.sheetnames if s.strip().upper() == target_sheet_name]
                    if not target_sheets:
                        target_sheets = [s for s in wb.sheetnames if target_sheet_name in s.strip().upper()]

                    for sname in target_sheets[:1]:
                        ws = wb[sname]
                        status_col = 11
                        edd_col = 12
                        act_col = 13
                        time_col = 14

                        ws.cell(row=1, column=status_col, value="Live Status").fill = header_fill
                        ws.cell(row=1, column=status_col).font = header_font
                        ws.cell(row=1, column=edd_col, value="EDD").fill = header_fill
                        ws.cell(row=1, column=edd_col).font = header_font
                        ws.cell(row=1, column=act_col, value="Latest Activity").fill = header_fill
                        ws.cell(row=1, column=act_col).font = header_font
                        ws.cell(row=1, column=time_col, value="Activity Date/Time").fill = header_fill
                        ws.cell(row=1, column=time_col).font = header_font

                        for row in ws.iter_rows(min_row=2):
                            row_idx = row[0].row
                            row_docket = None
                            docket_cell = None
                            for cell in row[:10]:
                                if cell.value:
                                    val_str = str(cell.value).strip()
                                    if c_type == 'trackon':
                                        if re.match(r'^[0-9]{10,12}$', val_str):
                                            row_docket = val_str
                                            docket_cell = cell
                                            break
                                    else:
                                        if re.match(r'^[A-Z]{1,4}[0-9]{6,12}$', val_str, re.I):
                                            row_docket = val_str.upper()
                                            docket_cell = cell
                                            break

                            if row_docket and row_docket in status_map:
                                info = status_map[row_docket]
                                c_status = ws.cell(row=row_idx, column=status_col, value=info.get('status', ''))
                                c_status.font = status_font
                                st_lower = str(info.get('status', '')).lower().strip()
                                cat_lower = str(info.get('category', '')).lower().strip()
                                is_delivered = (cat_lower == 'delivered' or
                                                st_lower == 'delivered' or
                                                st_lower.startswith('delivered') or
                                                ('delivered' in st_lower and not any(k in st_lower for k in ['undelivered', 'not delivered', 'failed', 'rto', 'return', 'cancel'])))
                                if is_delivered:
                                    c_status.fill = fill_delivered
                                    c_status.font = font_delivered
                                    # Highlight the docket number cell green only when delivered!
                                    if docket_cell is not None:
                                        docket_cell.fill = fill_delivered
                                        docket_cell.font = font_delivered
                                else:
                                    if docket_cell is not None:
                                        docket_cell.fill = PatternFill(fill_type=None)
                                        docket_cell.font = regular_font
                                    if 'out for delivery' in st_lower or cat_lower == 'out for delivery':
                                        c_status.fill = fill_ofd
                                    elif any(k in st_lower for k in ['undelivered', 'not delivered', 'failed', 'rto', 'return', 'cancel']) or cat_lower in ['issue/rto', 'not found']:
                                        c_status.fill = fill_issue
                                    else:
                                        c_status.fill = fill_transit

                                ws.cell(row=row_idx, column=edd_col, value=info.get('edd', '-')).font = regular_font
                                ws.cell(row=row_idx, column=act_col, value=info.get('latest_update', '-')).font = regular_font
                                ws.cell(row=row_idx, column=time_col, value=info.get('latest_time', '-')).font = regular_font

                out_path = os.path.join(UPLOAD_FOLDER, 'temp_merged_export.xlsx')
                wb.save(out_path)
                wb.close()
                print("Merge complete! Sending file...", flush=True)

                return send_file(
                    out_path,
                    as_attachment=True,
                    download_name="DAILY_ENTRY_2026_LIVE_STATUS.xlsx",
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            except Exception as e:
                print("Failed to merge into original:", e, flush=True)

    # Standalone summary workbook fallback (Fast, 0.05s)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Combined Tracking Status" if courier == 'both' else ("Trackon Tracking Status" if courier == 'trackon' else "DTDC Tracking Status")
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "S.No", "Courier", "Consignment / AWB No", "Reference No", "Category", 
        "Live Status", "EDD", "Origin", "Destination", "Latest Activity", "Activity Date/Time"
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
    ws.row_dimensions[1].height = 26

    cat_colors = {
        'Delivered': PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        'Out for Delivery': PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
        'In Transit': PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),
        'Booked': PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
        'Issue/RTO': PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        'Not Found': PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    }

    for idx, r in enumerate(results, start=1):
        row_num = idx + 1
        cat = r.get('category', 'In Transit')
        c_label = "Trackon" if r.get('courier') == 'trackon' else "DTDC"
        row_data = [
            idx, c_label, r.get('awb', ''), r.get('reference_no', '-'), cat,
            r.get('status', ''), r.get('edd', '-'), r.get('origin', '-'),
            r.get('destination', '-'), r.get('latest_update', '-'), r.get('latest_time', '-')
        ]
        ws.append(row_data)
        ws.row_dimensions[row_num].height = 20
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 2, 4, 5, 7]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align
            if col_idx == 5 and cat in cat_colors:
                cell.fill = cat_colors[cat]
                if cat == 'Delivered':
                    cell.font = Font(name="Calibri", size=10, bold=True, color="006100")
            # Highlight docket number in green when delivered
            if col_idx == 3 and cat == 'Delivered':
                cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                cell.font = Font(name="Calibri", size=10, bold=True, color="006100")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    summary_name = "Combined_Tracking_Summary.xlsx" if courier == 'both' else ("Trackon_Tracking_Summary.xlsx" if courier == 'trackon' else "DTDC_Tracking_Summary.xlsx")
    return send_file(
        buf,
        as_attachment=True,
        download_name=summary_name,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print(f"Starting DTDC Bulk Tracker Dashboard on http://0.0.0.0:{port} ...")
    app.run(host='0.0.0.0', port=port, debug=False)

